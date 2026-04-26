# Standard Library
import io
from pathlib import Path

# Private Library
from tests.fixtures._image import minimal_png


def create_sample_docx(dest: Path) -> None:
    """Create a minimal structured .docx at *dest* using python-docx.

    Structure mirrors sample.pdf and sample.md:
      H1 "Sample Document", H2 "Introduction" / "Data Overview" / "Conclusion",
      one embedded image and one GFM-style table under Data Overview.
    """
    from docx import Document
    from docx.shared import Inches

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    doc.add_heading("Sample Document", level=1)

    doc.add_heading("Introduction", level=2)
    doc.add_paragraph("This is the introduction paragraph.")

    doc.add_heading("Data Overview", level=2)
    doc.add_paragraph("This section covers data.")

    # Embed a small image
    doc.add_picture(io.BytesIO(minimal_png()), width=Inches(0.5))

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Alpha"
    table.rows[1].cells[1].text = "1"
    table.rows[2].cells[0].text = "Beta"
    table.rows[2].cells[1].text = "2"

    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph("Final remarks go here.")

    doc.save(str(dest))
