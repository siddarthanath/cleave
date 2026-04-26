# ───────────────────────────────────────────────────── Imports ────────────────────────────────────────────────────── #

# Standard Library
import base64
import re
from typing import List

# Third Party Library
try:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph
except ImportError as e:
    raise ImportError(
        "DocxParser requires python-docx. Install it with: pip install cleave[docx]"
    ) from e

# Private Library
from cleave.parsers.base import BaseModeParser
from cleave.schemas import (
    ContentBlock,
    ContentType,
    Document,
    DocumentPage,
    TreeNode,
)

# ────────────────────────────────────────────────────── Code ──────────────────────────────────────────────────────── #

# Matches "Heading 1", "Heading 2", … up to level 9
_HEADING_RE = re.compile(r"^Heading (\d)$", re.IGNORECASE)


class DocxParser(BaseModeParser):
    """Parser for .docx documents supporting flat and tree output modes.

    Flat mode  — all content on a single virtual page (page_number=None).
                 Paragraphs → text ContentBlocks, tables → markdown ContentBlocks,
                 inline images → base64 ContentBlocks.

    Tree mode  — TreeNode hierarchy built from Word paragraph styles.
                 'Heading N' styles define depth; body paragraphs become children
                 of the nearest ancestor heading. Tables and images are atomic
                 leaf nodes attached to the current heading scope.
    """

# ╔════════════════════════════════════════════════════════════════════════════════════╗
# ║                                    FLAT PATH                                       ║
# ║                        (Sequential Page-Block Extraction)                          ║
# ╚════════════════════════════════════════════════════════════════════════════════════╝

    def _parse_flat(self) -> Document:
        """Extract all content into a single virtual DocumentPage."""
        doc = DocxDocument(self.source.location)
        blocks: List[ContentBlock] = []
        position = 0

        for element in self._iter_block_elements(doc):
            if isinstance(element, Paragraph):
                text = element.text.strip()
                if text:
                    blocks.append(ContentBlock(
                        type=ContentType.text,
                        content=text,
                        position=position,
                    ))
                    position += 1
                for b64 in self._extract_paragraph_images(element):
                    blocks.append(ContentBlock(
                        type=ContentType.image,
                        content=b64,
                        position=position,
                    ))
                    position += 1

            elif isinstance(element, DocxTable):
                md = self._table_to_markdown(element)
                if md:
                    blocks.append(ContentBlock(
                        type=ContentType.table,
                        content=md,
                        position=position,
                    ))
                    position += 1

        if not blocks:
            raise ValueError(f"No extractable content found in '{self.source.location}'")

        page = DocumentPage(page_number=None, blocks=blocks)
        return Document(source=self.source, pages=[page], total_pages=1)

# ╔════════════════════════════════════════════════════════════════════════════════════╗
# ║                                    TREE PATH                                       ║
# ║                        (Heading-Scoped Hierarchy Building)                         ║
# ╚════════════════════════════════════════════════════════════════════════════════════╝

    def _parse_tree(self) -> Document:
        """Build a TreeNode hierarchy using Word paragraph heading styles."""
        doc = DocxDocument(self.source.location)

        root = TreeNode(
            content_type=ContentType.text,
            content="",
            metadata={"role": "root"},
        )
        # Stack entries: (heading_level, node) — level 0 = root, never popped
        stack: List[tuple] = [(0, root)]

        for element in self._iter_block_elements(doc):
            if isinstance(element, Paragraph):
                self._process_paragraph_tree(element, stack)
            elif isinstance(element, DocxTable):
                md = self._table_to_markdown(element)
                if md:
                    node = TreeNode(
                        content_type=ContentType.table,
                        content=md,
                        metadata={},
                    )
                    stack[-1][1].children.append(node)

        return Document(source=self.source, root=root)

    def _process_paragraph_tree(self, para: "Paragraph", stack: list) -> None:
        """Classify a paragraph as heading or body text and attach it to the tree.

        Heading paragraphs adjust the stack depth; body paragraphs are attached
        as children of the current top-of-stack node.

        Args:
            para: The paragraph element.
            stack: Mutable list of (level, TreeNode) tuples.
        """
        text = para.text.strip()
        style_name = para.style.name if para.style else ""
        match = _HEADING_RE.match(style_name)

        if match:
            level = int(match.group(1))
            while len(stack) > 1 and stack[-1][0] >= level:
                stack.pop()
            node = TreeNode(
                content_type=ContentType.text,
                content=text,
                metadata={"role": "heading", "level": level},
            )
            stack[-1][1].children.append(node)
            stack.append((level, node))
        else:
            for b64 in self._extract_paragraph_images(para):
                img_node = TreeNode(
                    content_type=ContentType.image,
                    content=b64,
                    metadata={},
                )
                stack[-1][1].children.append(img_node)
            if text:
                node = TreeNode(
                    content_type=ContentType.text,
                    content=text,
                    metadata={"role": "paragraph"},
                )
                stack[-1][1].children.append(node)

    @staticmethod
    def _iter_block_elements(doc: "DocxDocument"):
        """Yield top-level Paragraph and Table elements in document order.

        python-docx's `doc.paragraphs` and `doc.tables` are separate lists and do
        not preserve interleaved order — walking the XML body directly does.

        Args:
            doc: Open python-docx Document.

        Yields:
            Paragraph | Table elements in document order.
        """
        from docx.oxml.ns import qn as _qn
        body = doc.element.body
        para_tag = _qn("w:p")
        tbl_tag = _qn("w:tbl")
        for child in body:
            if child.tag == para_tag:
                yield Paragraph(child, doc)
            elif child.tag == tbl_tag:
                yield DocxTable(child, doc)

    @staticmethod
    def _extract_paragraph_images(para: "Paragraph") -> List[str]:
        """Extract base64-encoded images embedded as inline drawings in a paragraph.

        Args:
            para: The paragraph to search.

        Returns:
            Base64-encoded image strings, one per image found.
        """
        images: List[str] = []
        try:
            blip_tag = qn("a:blip")
            embed_attr = qn("r:embed")
            for drawing in para._element.iter():
                if drawing.tag == blip_tag:
                    rId = drawing.get(embed_attr)
                    if rId and rId in para.part.rels:
                        img_part = para.part.rels[rId].target_part
                        b64 = base64.b64encode(img_part.blob).decode("utf-8")
                        images.append(b64)
        except Exception:
            pass
        return images

    @staticmethod
    def _table_to_markdown(table: "DocxTable") -> str:
        """Convert a python-docx Table to a GitHub-Flavoured Markdown table string.

        Args:
            table: The table element.

        Returns:
            Markdown table string, or "" if the table has no rows.
        """
        rows = table.rows
        if not rows:
            return ""
        lines: List[str] = []
        for i, row in enumerate(rows):
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join("---" for _ in cells) + " |")
        return "\n".join(lines)
